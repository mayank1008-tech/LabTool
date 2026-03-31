from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from app.models.schemas import TemplateProfile, TextStyle


def build_docx(
    profile: TemplateProfile,
    experiment_number: str | None,
    aim: str,
    source_code: str,
    output_text: str | None,
) -> Document:
    doc = Document()

    _apply_page_margins(doc, profile)

    heading_label = profile.experiment_number_label
    if experiment_number:
        heading_label = f"{profile.experiment_number_label} {experiment_number}"

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(heading_label)
    _apply_style(run, profile.heading_style)

    _add_section(doc, profile.aim_section.label, aim, profile.aim_section)

    _add_label_paragraph(doc, profile.source_code_section.label, profile.source_code_section)
    code_para = doc.add_paragraph()
    code_run = code_para.add_run(source_code)
    _apply_style(code_run, profile.source_code_section.body_style)

    _add_label_paragraph(doc, profile.output_section.label, profile.output_section)
    out_para = doc.add_paragraph()
    out_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    if output_text:
        out_run = out_para.add_run(output_text)
        _apply_style(out_run, profile.output_section.body_style)
    else:
        placeholder_text = f"\n\n{profile.output_placeholder}\n\n"
        out_run = out_para.add_run(placeholder_text)
        out_run.font.color.rgb = RGBColor(180, 180, 180)
        out_run.italic = True

    return doc


def _add_section(doc: Document, label: str, body: str, section_cfg) -> None:
    p = doc.add_paragraph()
    label_run = p.add_run(label)
    _apply_style(label_run, section_cfg.label_style)
    body_run = p.add_run(body)
    _apply_style(body_run, section_cfg.body_style)


def _add_label_paragraph(doc: Document, label: str, section_cfg) -> None:
    p = doc.add_paragraph()
    label_run = p.add_run(label)
    _apply_style(label_run, section_cfg.label_style)


def _apply_style(run, style: TextStyle) -> None:
    run.font.name = style.font_family
    run.font.size = Pt(style.font_size)
    run.bold = style.bold
    run.italic = style.italic
    run.font.underline = style.underline


def _apply_page_margins(doc: Document, profile: TemplateProfile) -> None:
    layout = profile.page_layout
    for section in doc.sections:
        section.top_margin = Cm(layout.margin_top_cm)
        section.bottom_margin = Cm(layout.margin_bottom_cm)
        section.left_margin = Cm(layout.margin_left_cm)
        section.right_margin = Cm(layout.margin_right_cm)
